#!/usr/bin/env python3

from tables import *
import os, re, sys
import logging
from collections import defaultdict, OrderedDict
import glob
import inspect, time
from openpyxl import load_workbook
import docx
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from docx.enum.text import WD_UNDERLINE
import regex
import subprocess

logging.basicConfig(format='[%(asctime)s,%(msecs)03d] %(message)s', level=logging.INFO, datefmt='%H:%M:%S',)

YDS = {"+":"又", "-":"白", "*":"俗", "/":"書","\\":"語","=":"文","?":"存疑", "@": "訓"}
IPA_SM = "(ʔ?[ʡˀʕʢˤbɓᵇɸβʙpmɰɱᵐfᶠʩɟdɗɖȡᶑᵈʣʤʥꭦðtʈȶᵗʦʧʨꭧθnŋᵑɲɳȵɴᶰlɬɭʟ𝼄ɮ𝼅ʪʫgɡɠᶢɢʛ𐞒kʞhħɦʱɧʰʜzʐʑʒʓcʗCçɕsʂşȿrɹɻɺ𝼈ɾɽʀʁʃʄʆjʝʲq𐞥xχˣvʋⱱɣwẘʍʷʎ𝼆∅ʼ'\\.·~ʘǀǃǁǂ𝼊\u0300-\u0362]*)"
IPA_PATTERN = re.compile(f"^{IPA_SM}([^\\d⁰¹²³⁴⁵⁶˩˨˧˦˥]+)?([\\d⁰¹²³⁴⁵⁶˩˨˧˦˥]+[a-z]?)?([\\+\\-=\\?\\*@])?$")

def getYDMark(py):
	return py[-1] if py[-1] in YDS else ""

def getYD(py):
	# if py[-1] in ("-", "="):
	# 	return ""
	return YDS.get(py[-1], "")

def getCompatibilityVariants():
	d = dict()
	fname = os.path.join(WORKSPACE, "tables", "data", "HZUnified.txt")
	if not os.path.exists(fname):
		fname = os.path.join(WORKSPACE, "..", "app/src/main/assets/opencc/HZUnified.txt")
	for 行 in open(fname, encoding="U8"):
		字, val = 行.rstrip().split("\t")
		d[字] = val
	return d

def maketrans(path):
	fname = os.path.join(WORKSPACE, "tables", "data", path)
	d = dict()
	for 行 in open(fname, encoding="U8"):
		字, val = 行.rstrip("\n").split("\t")
		d[字] = val
	return str.maketrans(d)

def getTsvName(fname, 頁名=""):
	name = os.path.basename(fname)
	if 頁名: 頁名 = "-" + 頁名
	name = re.sub(r" ?(\(\d{0,3}\))+$", "", name.rsplit(".", 1)[0]) + 頁名 + ".tsv"
	return os.path.join(PATH, SOURCE, name)

def isXlsx(fname):
	return fname.endswith(".xlsx") or fname.endswith(".xls")

def processFs(v):
	t = type(v)
	if t is float or t is int: return "%d" % v
	if v is None: return ""
	return str(v).strip().replace("\t", " ").replace("\n", " ")

def processXlsxFs(v):
	t = type(v)
	if t is float or t is int: return "%d" % v
	if t is bool: return str(v)
	if v is None: return ""
	if t is str: return str(v).strip().replace("\t", " ").replace("\n", " ")
	cells = []
	for i in v:
		if type(i) is str:
			cells.append(i.strip())
			continue
		if type(i) is int or type(i) is float:
			cells.append("%d" % i)
			continue
		text = i.text
		tag = ""
		if i.font.underline == "single":
			tag = "-"
		elif i.font.underline == "double":
			tag = "="
		if tag:
			text = "".join([j + tag for j in text])
		# elif i.font.bold:
		# 	text = f"**{text}**"
		elif i.font.color and i.font.color.rgb == "FF808080":
			text = f"`{text}`"
		if i.font.vertAlign == "subscript" or (i.font.size and i.font.size < 10.0):
			text = f"({text})"
		cells.append(text)
	return "".join(cells).replace(")(", "").strip().replace("\n", "\\n")

def getXlsxLines(xlsx, 頁名):
	wb = load_workbook(xlsx, data_only=True, rich_text=True)
	sheet = wb[頁名] if 頁名 else wb.active
	lines = list()
	for row in sheet.rows:
		列 = [processXlsxFs(j.value) for j in row[:50]]
		if any(列):
			行 = "\t".join(列) + "\n"
			lines.append(行)
	return lines

def xlsx2tsv(xlsx, 頁名):
	tsv = getTsvName(xlsx, 頁名)
	if not os.path.exists(xlsx): return
	if os.path.exists(tsv):
		xtime = os.path.getmtime(xlsx)
		ttime = os.path.getmtime(tsv)
		if ttime >= xtime: return
	lines = getXlsxLines(xlsx, 頁名)
	t = open(tsv, "w", encoding="U8", newline="\n")
	t.writelines(lines)
	t.close()

def run2text(run):
	if isinstance(run, docx.text.hyperlink.Hyperlink):
		return "".join(map(run2text, run.runs))
	tag = ""
	if run.font.underline == WD_UNDERLINE.SINGLE:
		tag = "-"
	elif run.font.underline == WD_UNDERLINE.DOUBLE:
		tag = "="
	elif run.font.underline == WD_UNDERLINE.WAVY:
		tag = chr(0x1AB6)
	elif run._r.xpath("*/w:em[@w:val='dot']"):
		tag = chr(0x0323)
	text = run.text
	if tag:
		text = "".join([i + tag for i in text])
	if run.font.subscript or (run.font.size and run.font.size < Pt(9)):
		if text.startswith("{") and text.endswith("}"):
			pass
		# elif text.startswith("[") and text.endswith("]"):
		# 	pass
		else:
			text = f"{{{text}}}".replace("-", "").replace("=", "")
	return text

def isDocx(fname):
	return fname.endswith(".docx") or fname.endswith(".doc")
	
def docx2tsv(fname):
	tsv = getTsvName(fname)
	if not os.path.exists(fname): return
	if os.path.exists(tsv):
		xtime = os.path.getmtime(fname)
		ttime = os.path.getmtime(tsv)
		if ttime >= xtime: return
	lines = []
	Doc = Document(fname)
	for each in Doc._body._element:
		if isinstance(each, docx.oxml.table.CT_Tbl):
			t = Table(each, Doc)
			for row in t.rows:
				行 = ""
				cells = row.cells
				for i, cell in enumerate(cells):
					if cell in cells[:i]: continue
					for p in cell.paragraphs:
						行 += "".join(map(run2text, p.iter_inner_content())).replace("\t", "").replace("\n", "")
					行 += "\t"
				lines.append(行.replace("}~", "~}").replace("~{", "{~").replace("}{", "").replace("[}", "}[").replace("{h}", "h").rstrip())
		elif isinstance(each, docx.oxml.text.paragraph.CT_P):
			element = Paragraph(each, Doc)
			行 = "".join(map(run2text, element.iter_inner_content())).replace("}~", "~}").replace("~{", "{~").replace("}{", "").replace("[}", "}[").replace("{h}", "h")
			lines.append(行)
	行 = "\n".join(lines).replace("}\n{", "").replace("\n}", "}\n")
	t = open(tsv, "w", encoding="U8", newline="\n")
	t.write(行)
	t.close()

def ybKey(x):
	if "\t" not in x:
		return x[-1]
	音, 註 = x.split("\t", 1)
	if 註: 註 = 註[0]
	return 註 + 音[-1]

class 表:
	_time = os.path.getmtime(__file__)
	文件名 = None
	頁名 = ""
	_sep = None
	顏色 = "#1E90FF"
	全稱 = ""
	簡稱 = ""
	說明 = ""
	網站 = ""
	網址 = ""
	字書 = False

	註序 = False
	補丁 = None
	kCompatibilityVariants = getCompatibilityVariants()
	kPUAs = maketrans("私用字.tsv")
	kShapes = maketrans("正形.tsv")
	kIPAs = maketrans("正音.tsv")
	kTones = maketrans("正調.tsv")
	simplified = 1
	爲音 = True
	列序 = None
	音節典 = defaultdict(set)
	音表 = OrderedDict()
	聲韻典 = defaultdict(set)
	d = defaultdict(list)
	__mod = None
	誤 = []
	音集 = set()
	不計入調 = set()
	調號 = "˩˨˧˦˥⁰¹²³⁴⁵⁶⁷⁸⁹⓪①②③④⑤⑥⑦⑧⑨ⓐⓑⓒⓓ"
	韻母集 = set()
	聲母集 = set()
	聲調典 = defaultdict(set)

	def __init__(自):
		自.誤.clear()
		自.音集.clear()

	def setmod(自, mod):
		自.__mod = mod

	def __str__(自):
		if 自.__mod: return 自.__mod
		return 自.__module__.split(".")[-1]

	def find(自, name):
		if os.sep not in name and (isXlsx(name) or isDocx(name)):
			name = 自.toolname(name)
			if g := 自.find(name): return g
		if os.sep not in name:
			name = 自.全路徑(name)
		if g := glob.glob(name): return g
		if g := glob.glob(glob.escape(name)): return g
		if g := glob.glob(re.sub(".([^.]+)$", "([0-9]).\\1", name)): return g
		if g := glob.glob(re.sub(".([^.]+)$", "([0-9][0-9]).\\1", name)): return g
		if g := glob.glob(re.sub(".([^.]+)$", " ([0-9]).\\1", name)): return g
		if g := glob.glob(re.sub(".([^.]+)$", " ([0-9][0-9]).\\1", name)): return g
		if isXlsx(name) or isDocx(name):
			自.文件名 = getTsvName(自.文件名, 自.頁名)
			return 自.find(自.文件名)
		return

	@property
	def spath(自):
		if 自.文件名 and "/" in 自.文件名:
			自.文件名, 自.頁名 = 自.文件名.rsplit("/", 1)
		sname = 自.文件名
		if not 自.簡稱: 自.簡稱 = 自.info["簡稱"]
		if not 自.簡稱: 自.簡稱 = str(自)
		if not sname: sname = f"{自.簡稱}.tsv"
		g = 自.find(sname)
		if not g:
			logging.error(f"\t\t未找到 {sname}")
			自.文件名 = None
			return
		if len(g) != 1:
			logging.error(f"\t\t找到多个 {sname}：{g}")
			return
		sname = g[0]
		自.文件名 = os.path.basename(sname)
		if isXlsx(sname):
			xlsx2tsv(sname, 自.頁名)
			sname = getTsvName(sname, 自.頁名)
		elif isDocx(sname):
			docx2tsv(sname)
			sname = getTsvName(sname)
		return sname

	def toolname(自, name):
		name = os.path.basename(name)
		return os.path.join(PATH, "..", name)

	def 全路徑(自, name):
		name = os.path.basename(name)
		return os.path.join(PATH, SOURCE, name)

	@property
	def tpath(自):
		tpath = os.path.join(PATH, TARGET, 自.簡稱)
		if not tpath.endswith(".tsv"): tpath += ".tsv"
		return tpath

	def normS(自, s, rep="｛\\1｝"):
		s = regex.sub(r"\(((?>[^\(\)]+|(?R))*)\)", rep, s)
		return s

	def normM(自, s, rep="｛\\1｝"):
		s = regex.sub(r"\[((?>[^\[\]]+|(?R))*)\]", rep, s)
		return s

	def normG(自, s, rep="｛\\1｝"):
		s = regex.sub(r"\{((?>[^\{\}]+|(?R))*)\}", rep, s)
		return s

	def 過時(自):
		classfile = inspect.getfile(自.__class__)
		classtime = os.path.getmtime(classfile)
		varianttime = os.path.getmtime(VARIANT_FILE)
		if classtime < varianttime:
			classtime = varianttime
		spath = 自.spath
		if not spath or not os.path.exists(spath):
			return False
		if "版本" in 自.info and not 自.info["版本"]:
			result = subprocess.run(["git", "log", "-1", "--format=%cd", "--date=short", spath], stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True)
			if result.returncode == 0:
				版本 = result.stdout.strip()
			else:
				版本 = time.strftime("%Y-%m-%d", time.localtime(os.path.getmtime(spath)))
			自.info["版本"] = 版本
		if os.path.exists(自.tpath):
			ftime = os.path.getmtime(spath)
			ttime = os.path.getmtime(自.tpath)
			if ttime < 自._time: return True
			if ttime < classtime: return True
			return ttime < ftime
		return True

	def 修訂(自, d):
		if not 自.補丁: return
		for 字, 音 in 自.補丁.items():
			if not 音:
				del d[字]
				continue
			d[字] = 音.split(",")

	def 無調(自):
		return 自.簡稱.endswith("上古") or 自.簡稱.endswith("朝鮮") or 自.簡稱.startswith("日語") or 自.簡稱 in ("1851寧波", "1884新甯", "1890新會", "党項")
	
	def 無q聲(自):
		return 自.簡稱 not in ("盛唐", "榕江侗上古借詞", "榕江侗中古借詞") and not 自.文件名.startswith("白語")

	def 無音(自, 音):
		音 = 音.strip(" /-—－")
		return 自.爲語() and 自.爲音 and (音.isdigit() or 音 == "")
	
	def 校音(自, 音):
		名 = 自.簡稱
		if 名 in ("南寧", "南寧亭子"): 音 = 音.replace("-", "/")
		if 有字(音):
			自.誤.append(f"[{音}]音節含有漢字")
			print(f"{自.簡稱} \"{自.spath}\" [{音}]音節含有漢字")
		for i in 音.replace("`", "").replace("*", "").replace("~", "/").split("/"):
			if m := IPA_PATTERN.match(i):
				ym = m.group(2)
				if not ym: ym = ""
				ym = ym.rstrip("-=+*@?")
				自.韻母集.add(ym)
				sm = m.group(1)
				if not sm: sm = ""
				if len(sm) >= 2 and sm[-1] in "\u0329\u030D":
					sm = sm[:-2]
					if not ym: ym = sm[-2:]
				自.聲母集.add(sm)
				if m.group(3): 自.聲調典[自.簡稱].add(m.group(3))
			else:
				自.誤.append(f"[{音}]音節含有非法字符")
				print(f"{自.簡稱} \"{自.spath}\" [{音}]音節含有非法字符")

	def _正音(自, 音):
		if 自.爲語() and 自.爲音:
			if 音 in ("", "-", "—", "－", "一", "/", "误", "∅"):
				return ""
			音 = 音.strip("[] ")
			音 = 音.replace("Ǿ", "Ǿ").replace("Ǿ", "").lstrip("∅︀∅Ø〇0").replace("零", "").replace("◌", "")
			if 自.無q聲(): 音 = 音.lstrip("q")
			音 = re.sub("^[I1]", "l", 音)
			音 = 音.lower().translate(自.kIPAs).replace("tc", "tɕ").translate(自.kTones)
			音 = re.sub("([ʂʐ]ʰ?)ʮ", "\\1ʯ", 音)
			音 = re.sub("([sz]ʰ?)ʯ", "\\1ʮ", 音)
			音 = re.sub("([ʂʐ]ʰ?)ɿ", "\\1ʅ", 音)
			音 = re.sub("([sz]ʰ?)ʅ", "\\1ɿ", 音)
			音 = re.sub(r"[\[\]\{\} ]","", 音)
			#(i)e-> ie/e
			if "(" in 音 and ")" in 音:
				音甲 = 音.replace("(", "").replace(")", "")
				音乙 = re.sub(r"\(.*?\)", "", 音)
				音 = f"{音甲}/{音乙}"
			音 = re.sub("([^/])h", "\\1ʰ", 音)
			音 = re.sub(fr"^{IPA_SM}([mnvʋzʑrɹɻlɭβŋȵʐɱʒ])(\d+)$", "\\1\\2\u0329\\3", 音)
			音 = re.sub("([mnvʋzʑrɹɻlɭβ])([\u0329\u030Dˌˈ]+)", "\\1\u0329", 音)
			音 = re.sub("([ŋȵʐɱʒ])([\u0329\u030Dˌˈ]+)", "\\1\u030D", 音)
			音 = re.sub("^([^*])\\1([^\u0303])", "\\1\\2", 音)
			音 = re.sub("([-=])(\\d+)$", "\\2\\1", 音)
			自.校音(音)
			if 自.無調():
				音 = 音.rstrip("0123456789")
		return 音

	def 正音(自, 音, 檢查=False):
		if 自.無音(音): return ""
		音 = 自._正音(音)
		if not 檢查: return 音
		if "\t" in 音:
			自.誤.append(f"[{音}]音節含TAB字符")
			音 = 音.replace("\t", "")
		if 音 not in 自.音集:
			自.音集.add(音)
		elif 自.簡稱 not in ("長沙星沙", "長沙金井", "雙牌打鼓坪", "湘劇", "蘇州評彈", "溧陽河南話", "南京", "新洲", "1800長沙", "鹽城西鄉") and 自.文件名 not in ("临沧地区汉语方言志.tsv",):
			自.誤.append(f"[{音}]音節重複")
		return 音

	def 檢查同音字(自):
		return 自.分區 and 自.簡稱 not in ("普通話",) and not 自.分區.startswith("歷史音") and not 自.分區.startswith("域外方音")

	def 爲方言(自):
		return 自.簡稱 in ("老國音","党項") or (自.爲語() and not 自.分區.startswith("歷史音"))

	def 去括號(自, 註):
		if not 註: return ""
		if 註[0] == '{' and 註[-1] == '}' and "{" not in 註[1:-1]:
			註 = 註[1:-1]
		elif 註[0] == '(' and 註[-1] == ')' and "(" not in 註[1:-1]:
			註 = 註[1:-1]
		return 註.strip()

	def 正註(自, 註):
		if not 註: return ""
		註 = 自.去括號(註)
		上 = ""
		果 = list()
		for 字 in 註:
			if 爲字(字) or 字 == "~":
				if 上: 果.append(上)
				上 = ""
				果.append(字)
			else:
				上 += 字
		if 上: 果.append(上)
		return re.sub(r" ?([,:;?!()]) ?", "\\1", " ".join(果).replace("   ", "  "))

	def 合註(自, 註):
		return 註.replace("  ", "　").replace(" ", "").replace("　", " ")

	def 正部件(自, 註):
		if not 註: return ""
		上 = ""
		果 = list()
		for 部件 in 註:
			if len(部件.encode()) > 1:
				if 上: 果.append(上)
				上 = ""
				果.append(部件)
			else:
				上 += 部件
		if 上: 果.append(上)
		return " ".join(果)

	def 寫(自, d):
		自.修訂(d)
		t = open(自.tpath, "w", encoding="U8", newline="\n")
		print(f"#漢字\t音標\t解釋", file=t)
		for 字 in sorted(d.keys()):
			pys = d[字]
			字 = 自.kCompatibilityVariants.get(字, 字)
			if 自.爲方言() and 自.simplified:
				字 = s2t(字, 自.simplified)
			if not 爲字(字):
				if 自.爲方言():
					自.誤.append(f"【{字}】({','.join([i.strip() for i in pys])})不是漢字")
				continue
			if 自.註序:
				pys = sorted(pys,key=ybKey)
			for py in pys:
				if "\t" in py:
					音, 註 = py.split("\t", 1)
					註 = 註.strip()
				else:
					音, 註 = py, ""
				音 = 自.正音(音)
				if not 音: continue
				if 字 == "□" and not 註:
					自.誤.append(f"【□】({音})無註釋")
				音 = f"{音}\t{註}"
				print(f"{字}\t{音}", file=t)
		t.close()

	@property
	def 分區(自):
		return 自.info["地圖集二分區"]

	def 爲語(自):
		return 自.分區 != None

	@property
	def 字數(自):
		return len(自.d) + 自.框數 - (1 if 自.框數 > 0 else 0)
	
	@property
	def 框數(自):
		數 = len(自.d.get("□", []))
		if 自.爲語(): return 數
		return 1 if 數 > 0 else 0

	@property
	def 音節數(自):
		return len(自.音節典)

	@property
	def 聲韻數(自):
		return len(自.聲韻典)

	def 讀(自, 更新=False):
		自.音表.clear()
		自.音節典.clear()
		自.聲韻典.clear()
		自.d.clear()
		if (自.過時() or 更新) and 自.spath: 自.更新()
		if not 自.tpath or not os.path.exists(自.tpath): return
		for 行 in open(自.tpath,encoding="U8"):
			行 = 行.strip()
			if 行.startswith("#"): continue
			if "\t" not in 行: continue
			字, py = 行.split("\t", 1)
			if 自.爲語():
				註 = ""
				if "\t" in py: py, 註 = py.split("\t", 1)
				if 註 and 自.爲語():
					註 = 自.正註(註)
				try:
					異讀 = getYD(py)
				except:
					print("\t\t\t", 自.簡稱, py, 註)
					sys.exit(1)
				if 異讀 and py.count("*") <= 1:
					註 = f"({異讀}){註}"
					py = py[:-1]
				if re.match(r"^\([^()]*?\)$", 註):
					註 = 註[1:-1]
				音 = re.sub(r"\(.*?\)","",py).strip(" _`*")
				音 = 音.split("/", 1)[0]
				if "-" not in 音.rstrip("+-*/=?@\\"):
					繁註 = s2t(註.replace(" ", ""))
					if "兒化" not in 繁註 and "連讀" not in 繁註 and "語流" not in 繁註 and "變調" not in 繁註 and "合音" not in 繁註:
						音乙 = 音.rstrip("+-*/=?@\\")
						聲韻, 調 = 自.分音(音乙)
						if 自.無調() or (調 and not 調.startswith("0") and 調 not in 自.不計入調):
							自.音節典[音乙].add(字)
							自.聲韻典[聲韻].add(字)
				if 註:
					py += "{%s}" % 註
			else:
				if 自.字書:
					sep = "▲" if 自.簡稱 == "匯纂" else "\t"
					py2, 註 = py.split(sep, 1)
					py = ("\n\n" if 自.d[字] else "") + py2 + sep + 自.正註(註)
				elif 自.簡稱 in ("部件檢索","字形描述"):
					py = 自.正部件(py)
				py = py.replace("\t", "\n")
			if py not in 自.d[字]:
				自.d[字].append(py)

	def 加載(自, dicts, 更新=False):
		自.讀(更新)
		if not 自.d: return
		for 字, 音集 in 自.d.items():
			if 字 not in dicts:
				dicts[字] = {"漢字": 字}
			dicts[字][自.簡稱] = "\t".join(音集)

	def 加載條目(自, items, 更新=False):
		自.讀(更新)
		if not 自.d: return
		d = defaultdict(list)
		for 字, 音集 in 自.d.items():
			for 音 in 音集:
				d[音].append(字)
		for 音, 字組 in d.items():
			註 = re.sub(r"\{(.+?)\}$", "\t\\1", 音)
			if "\t" not in 註:
				讀音, 註釋 = 註, ""
				items.append((" ".join(字組), 自.簡稱, 讀音, 註釋))
			else:
				讀音, 註釋 = 註.split("\t", 1)
				for 字 in 字組:
					註釋乙 = 註釋
					if 自.爲語() and 自.爲音 and "~" in 註釋 and 字 != "□":
						註釋乙 = 註釋.replace("~", f"*{字}*")
						註釋乙 = re.sub(r"(\*) ([^* ])", "\\1\\2", 註釋乙)
						註釋乙 = re.sub(r"([^* ]) (\*)", "\\1\\2", 註釋乙)
					items.append((字, 自.簡稱, 讀音, 註釋乙))

	def 存(自, output):
		t = open(output, "w", encoding="U8", newline="\n")
		for 字, 音集 in 自.d.items():
			for 音 in 音集:
				註 = 自.合註(re.sub(r"\{([^{}]*?)\}$", "\t\\1", 音))
				t.write(f'{字}\t{註}\n')
		t.close()

	def 析(自, 列):
		return tuple(列[:3])

	def 統調(自, m):
		return m.group(0).translate(自.kTones)

	def 統(自, 行):
		行 = 行.rstrip('\n')
		行 = 行.translate(自.kPUAs).translate(自.kShapes)
		if not 自.爲方言(): return 行
		行 = 行.replace("...", "⋯").replace("ʔb", "ɓ").replace("ʔd", "ɗ")
		行 = re.sub(fr"\[[{自.調號}]+\]", 自.統調, 行)
		return 行
	
	@property
	def sep(自):
		if 自._sep: return 自._sep
		sep = "\t"
		spath = 自.spath
		if spath.endswith(".csv"): sep = ","
		elif spath.endswith(".tsv"): sep = "\t"
		elif spath.endswith(".txt"): sep = " "
		return sep

	def 更新(自):
		d = defaultdict(list)
		sep = 自.sep
		skip = 自.info.get("跳過行數", 0)
		lineno = 0
		files = [自.spath]
		for spath in files:
			for 行 in open(自.全路徑(spath),encoding="U8"):
				lineno += 1
				if lineno <= skip: continue
				行 = 自.統(行)
				if 行.startswith('#') : continue
				列 = [i.strip() for i in 行.split(sep)]
				entries = 自.析(列)
				if not entries: continue
				if type(entries) is tuple: entries = [entries]
				for 列 in entries:
					if len(列) <= 1: continue
					if len(列) >= 2:
						字, 音 = 列[:2]
						js = "\t".join(列[2:])
					if not 字 or len(字) != 1: continue
					if not 音: continue
					if 自.爲方言():
						if 爲字(音[0]): continue
					p = f"{音}\t{js}"
					p = p.strip()
					if p not in d[字]:
						d[字].append(p)
		自.寫(d)

	def 分音(自, 音):
		if not 音: return "",""
		音 = 音.translate(自.kTones)
		聲韻 = re.split(r"\d", 音, maxsplit=1)[0]
		調 = 音[len(聲韻):]
		return 聲韻,調

	def 轉調類(自, 音):
		音 = 音.strip().lstrip("0")
		if not 音: return 音
		異讀 = getYDMark(音)
		if 異讀:
			音 = 音[:-1]
		if re.findall(r"/[^\d]", 音):
			return "/".join(map(自.轉調類, re.split("/(?=[^\\d])", 音)))
		if "-" in 音:
			return "-".join(map(自.轉調類, 音.split("-")))
		聲韻,調值 = 自.分音(音)
		if not 調值: return 聲韻
		調類 = 自.僅轉調類(調值, 聲韻)
		return 聲韻 + 調類 + 異讀

	def 僅轉調類(自, 調值, 聲韻=""):
		調類 = ""
		if 調值 not in 自.調典:
			if 調值 == "0":
				調類 = 調值
			elif len(調值) == 1:
				調值 = 調值 * 2
				if 調值 in 自.調典:
					調類 = 自.調典[調值]
			elif len(調值) == 2 and 調值[0] == 調值[1]:
				調值 = 調值[0]
				if 調值 in 自.調典:
					調類 = 自.調典[調值]
			else:
				調類 = ""
		else:
			調類 = 自.調典[調值]
		if 聲韻 and 聲韻[-1] in "ptkʔ̚" and 調值 + "0" in 自.調典:
			調類 = 自.調典[調值 + "0"]
		return 調類
	
	def 增加調類(自, 行):
		字表使用調值 = 自.info.get("字表使用調值", False)
		調序 = list(自.調典.keys() if 字表使用調值 else 自.調典.values())
		調數 = len(調序)
		行 = "\t".join((f"[{調序[序-1]}]" if 0 < 序 <= 調數 else "") + 項 for 序,項 in enumerate(行.split("\t")))
		return 行
